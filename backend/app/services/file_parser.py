"""
Structure-aware document parsing.

Instead of flattening a PDF/DOCX into one giant text blob, we extract a list
of "elements" that each carry the structural metadata needed for good
chunking and citation later on:

    { "type": "heading" | "text" | "table",
      "text": str,
      "page_number": int | None,
      "section": str | None }

PDF: PyMuPDF (fitz) gives us font-size per text span, which we use as a cheap
     heuristic for heading detection. pdfplumber is used specifically for
     table extraction (converted to markdown so structure survives).

DOCX: python-docx exposes paragraph "style" (Heading 1/2/3...) directly, and
      native table objects, which we also convert to markdown.
"""
from typing import List, Dict, Any
import statistics

import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument


def _table_to_markdown(rows: List[List[str]]) -> str:
    rows = [[(c or "").strip() for c in row] for row in rows if row]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    md = "| " + " | ".join(header) + " |\n"
    md += "| " + " | ".join(["---"] * len(header)) + " |\n"
    for row in body:
        row = row + [""] * (len(header) - len(row))  # pad short rows
        md += "| " + " | ".join(row[: len(header)]) + " |\n"
    return md


def parse_pdf(path: str) -> List[Dict[str, Any]]:
    elements: List[Dict[str, Any]] = []
    current_section = None
    page_section_map: Dict[int, Any] = {}

    # --- Pass 1: text + heading detection via font size, page by page ---
    doc = fitz.open(path)
    all_sizes = []
    for page in doc:
        for block in page.get_text("dict").get("blocks", []):
            if "lines" not in block:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if span.get("text", "").strip():
                        all_sizes.append(round(span["size"], 1))
    body_size = statistics.median(all_sizes) if all_sizes else 10.0
    heading_threshold = body_size * 1.15  # anything notably bigger than body text = heading

    for page_num, page in enumerate(doc, start=1):
        blocks = page.get_text("dict").get("blocks", [])
        for block in blocks:
            if "lines" not in block:
                continue
            block_text_parts = []
            max_size = 0.0
            for line in block.get("lines", []):
                line_text = "".join(span.get("text", "") for span in line.get("spans", []))
                for span in line.get("spans", []):
                    max_size = max(max_size, span.get("size", 0.0))
                if line_text.strip():
                    block_text_parts.append(line_text.strip())
            block_text = " ".join(block_text_parts).strip()
            if not block_text:
                continue

            is_heading = max_size >= heading_threshold and len(block_text) < 150
            if is_heading:
                current_section = block_text
                elements.append({
                    "type": "heading", "text": block_text,
                    "page_number": page_num, "section": current_section,
                })
            else:
                elements.append({
                    "type": "text", "text": block_text,
                    "page_number": page_num, "section": current_section,
                })
        page_section_map[page_num] = current_section
    doc.close()

    # --- Pass 2: tables via pdfplumber (with page numbers and page-accurate sections) ---
    try:
        with pdfplumber.open(path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                page_section = page_section_map.get(page_num)
                for table in page.extract_tables():
                    md = _table_to_markdown(table)
                    if md.strip():
                        elements.append({
                            "type": "table", "text": md,
                            "page_number": page_num, "section": page_section,
                        })
    except Exception:
        pass  # table extraction is best-effort; text extraction above still succeeds

    # Preserve page order for elements
    elements.sort(key=lambda e: (e.get("page_number") or 0))
    return elements


def parse_docx(path: str) -> List[Dict[str, Any]]:
    elements: List[Dict[str, Any]] = []
    current_section = None
    doc = DocxDocument(path)

    # python-docx doesn't give a single reading-order stream of paragraphs+tables
    # out of the box, so we walk the underlying XML body in document order.
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]

        if tag == "p":
            # Find the matching Paragraph object to read style + text cleanly
            para = next((p for p in doc.paragraphs if p._p is child), None)
            if para is None or not para.text.strip():
                continue
            style_name = (para.style.name or "").lower()
            if "heading" in style_name or "title" in style_name:
                current_section = para.text.strip()
                elements.append({
                    "type": "heading", "text": para.text.strip(),
                    "page_number": None, "section": current_section,
                })
            else:
                elements.append({
                    "type": "text", "text": para.text.strip(),
                    "page_number": None, "section": current_section,
                })

        elif tag == "tbl":
            table = next((t for t in doc.tables if t._tbl is child), None)
            if table is None:
                continue
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            md = _table_to_markdown(rows)
            if md.strip():
                elements.append({
                    "type": "table", "text": md,
                    "page_number": None, "section": current_section,
                })

    return elements


def parse_document(path: str, filetype: str) -> List[Dict[str, Any]]:
    if filetype == "pdf":
        return parse_pdf(path)
    elif filetype == "docx":
        return parse_docx(path)
    else:
        raise ValueError(f"Unsupported filetype: {filetype}")
